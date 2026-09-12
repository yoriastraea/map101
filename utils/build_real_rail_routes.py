# -*- coding: utf-8 -*-
"""Match recorded railway journeys to real OpenStreetMap railway geometry.

The expensive nationwide railway graph is cached beside the private database.
Subsequent runs only rematch the workbook and regenerate the small public files.
"""

from __future__ import annotations

import argparse
import csv
import heapq
import json
import math
import pickle
import sys
import time
from collections import Counter, defaultdict
from pathlib import Path

# Avoid incompatible optional pandas accelerators in some Anaconda installs.
sys.modules.setdefault("numexpr", None)
sys.modules.setdefault("bottleneck", None)

import pyogrio
from openpyxl import load_workbook


CACHE_VERSION = 3
EARTH_RADIUS_KM = 6371.0088
REFERENCE_DATA = Path(__file__).resolve().parents[1] / "reference-data"


def parse_args():
    parser = argparse.ArgumentParser()
    parser.add_argument("database", type=Path)
    parser.add_argument("output", type=Path)
    parser.add_argument("--rebuild-network", action="store_true")
    return parser.parse_args()


def haversine(a, b):
    lon1, lat1 = map(math.radians, a)
    lon2, lat2 = map(math.radians, b)
    dlon, dlat = lon2 - lon1, lat2 - lat1
    value = math.sin(dlat / 2) ** 2 + math.cos(lat1) * math.cos(lat2) * math.sin(dlon / 2) ** 2
    return 2 * EARTH_RADIUS_KM * math.asin(min(1, math.sqrt(value)))


def geometry_length(coords):
    return sum(haversine(a, b) for a, b in zip(coords, coords[1:]))


def stitch_route_segments(segments):
    """Join ordered station-to-station paths into one continuous route."""
    stitched = []
    for segment in segments:
        if len(segment) < 2:
            continue
        if not stitched:
            stitched.extend(segment)
            continue
        # Adjacent paths may snap to different rail nodes around the same
        # station. Keeping both endpoints adds the short missing connector.
        if stitched[-1] == segment[0]:
            stitched.extend(segment[1:])
        else:
            stitched.extend(segment)
    return stitched


def normalized_line_name(value):
    name = str(value or "").strip().replace(" ", "")
    for token in ("高速铁路", "客运专线", "城际铁路", "铁路", "高速线", "城际线", "线"):
        name = name.replace(token, "")
    return name


def is_major_line(label, visited_km):
    excluded = ("联络", "疏解", "动车", "站", "道", "存车", "走行", "牵出", "渡线", "外绕", "到达", "出发", "机务", "车辆", "专用")
    looks_like_line = label.endswith(("线", "铁路")) or "高铁" in label or "客专" in label
    return visited_km >= 5 and looks_like_line and not any(token in label for token in excluded)


def iter_lines(geometry):
    if geometry is None or geometry.is_empty:
        return
    if geometry.geom_type == "LineString":
        yield geometry
    elif geometry.geom_type == "MultiLineString":
        yield from geometry.geoms


def build_network(source, cache_path):
    started = time.time()
    print(f"Reading railway geometry: {source}")
    clause = "railway = " + repr("rail")
    rails = pyogrio.read_dataframe(
        source,
        where=clause,
        columns=["id", "name", "name_zh"],
    )

    nodes = []
    node_lookup = {}
    edges = []
    adjacency = []
    parts = []
    shared_endpoints = set()

    def node_id(coord):
        key = (round(float(coord[0]), 6), round(float(coord[1]), 6))
        existing = node_lookup.get(key)
        if existing is not None:
            return existing
        index = len(nodes)
        node_lookup[key] = index
        nodes.append(key)
        adjacency.append([])
        return index

    for row in rails.itertuples(index=False):
        for part_index, line in enumerate(iter_lines(row.geometry)):
            coords = [(round(float(x), 6), round(float(y), 6)) for x, y, *_ in line.coords]
            if len(coords) < 2:
                continue
            name = row.name_zh or row.name or ""
            parts.append((str(row.id), part_index, str(name), coords))
            shared_endpoints.add(coords[0])
            shared_endpoints.add(coords[-1])

    # OSM ways may connect to an interior coordinate of another way. Split at
    # every coordinate that is an endpoint elsewhere so the graph is routable.
    for osm_id, part_index, name, coords in parts:
        segment_start = 0
        segment_number = 0
        for index in range(1, len(coords)):
            should_split = index == len(coords) - 1 or coords[index] in shared_endpoints
            if not should_split:
                continue
            segment = coords[segment_start:index + 1]
            segment_start = index
            if len(segment) < 2:
                continue
            start, end = node_id(segment[0]), node_id(segment[-1])
            if start == end:
                continue
            edge_index = len(edges)
            length_km = geometry_length(segment)
            edges.append({
                "osm_id": f"{osm_id}:{part_index}:{segment_number}",
                "name": name,
                "line_key": normalized_line_name(name),
                "length_km": length_km,
                "coords": segment,
                "u": start,
                "v": end,
            })
            adjacency[start].append((end, edge_index))
            adjacency[end].append((start, edge_index))
            segment_number += 1

    del rails, parts

    payload = {
        "version": CACHE_VERSION,
        "source_size": source.stat().st_size,
        "source_mtime": source.stat().st_mtime_ns,
        "nodes": nodes,
        "edges": edges,
        "adjacency": adjacency,
    }
    cache_path.parent.mkdir(parents=True, exist_ok=True)
    with cache_path.open("wb") as file:
        pickle.dump(payload, file, protocol=pickle.HIGHEST_PROTOCOL)
    print(f"Cached {len(nodes):,} nodes and {len(edges):,} edges in {time.time() - started:.1f}s")
    return payload


def load_network(source, cache_path, rebuild=False):
    if cache_path.exists() and not rebuild:
        print(f"Loading cached railway graph: {cache_path}")
        with cache_path.open("rb") as file:
            payload = pickle.load(file)
        if (
            payload.get("version") == CACHE_VERSION
            and payload.get("source_size") == source.stat().st_size
            and payload.get("source_mtime") == source.stat().st_mtime_ns
        ):
            return payload
        print("Railway source changed; rebuilding cache")
    return build_network(source, cache_path)


def load_stations(path):
    stations = {}
    with path.open(encoding="utf-8", newline="") as file:
        for row in csv.DictReader(file):
            try:
                stations[str(row["车站"]).strip()] = {
                    "coord": (float(row["经度"]), float(row["纬度"])),
                    "province": str(row.get("省") or ""),
                    "city": str(row.get("市") or ""),
                }
            except (KeyError, TypeError, ValueError):
                continue
    return stations


def sheet_records(sheet):
    rows = sheet.iter_rows(values_only=True)
    headers = [str(value).strip() if value is not None else "" for value in next(rows)]
    for values in rows:
        yield {headers[index]: value for index, value in enumerate(values) if index < len(headers) and headers[index]}


def load_trips(workbook_path):
    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    details = {row.get("序号"): row for row in sheet_records(workbook["乘坐列表"]) if row.get("序号") is not None}
    line_by_id = {row.get("序号"): str(row.get("线路") or "").strip() for row in sheet_records(workbook["线路"])}

    route_sheet = workbook["1"]
    rows = route_sheet.iter_rows(values_only=True)
    next(rows)
    trips = []
    for values in rows:
        if not values or values[0] is None:
            continue
        trip_id = values[0]
        detail = details.get(trip_id, {})
        station_names = [str(value).strip() for value in values[5:] if value not in (None, "")]
        if not station_names:
            station_names = [detail.get("上车站"), detail.get("下车站")]
        trips.append({
            "trip_id": trip_id,
            "date": detail.get("日期") or values[1],
            "train": str(detail.get("车次") or values[2] or ""),
            "from": str(detail.get("上车站") or values[3] or ""),
            "to": str(detail.get("下车站") or values[4] or ""),
            "duration": detail.get("时长.1") or 0,
            "distance": detail.get("里程") or 0,
            "fromcity": str(detail.get("上车城市") or ""),
            "tocity": str(detail.get("下车城市") or ""),
            "line": line_by_id.get(trip_id, ""),
            "stations": station_names,
        })
    return trips


def astar(network, starts, goals, start_coord, goal_coord, line_hint=""):
    start_set, goal_set = set(starts), set(goals)
    if start_set & goal_set:
        return []
    nodes, edges, adjacency = network["nodes"], network["edges"], network["adjacency"]
    hint = normalized_line_name(line_hint)
    max_goal_offset = max(haversine(nodes[index], goal_coord) for index in goal_set)
    queue = []
    best = {}
    for start in start_set:
        snap_cost = haversine(start_coord, nodes[start])
        best[start] = snap_cost
        heuristic = max(0, haversine(nodes[start], goal_coord) - max_goal_offset)
        heapq.heappush(queue, (snap_cost + heuristic, snap_cost, start))
    parent = {}
    selected_goal = None
    selected_total = math.inf

    while queue:
        priority, cost, node = heapq.heappop(queue)
        if priority >= selected_total:
            break
        if cost != best.get(node):
            continue
        if node in goal_set and node in parent:
            total = cost + haversine(nodes[node], goal_coord)
            if total < selected_total:
                selected_total = total
                selected_goal = node

        for neighbor, edge_index in adjacency[node]:
            edge = edges[edge_index]
            penalty = 1.0
            if hint and edge["line_key"] and hint not in edge["line_key"] and edge["line_key"] not in hint:
                penalty = 1.35
            next_cost = cost + edge["length_km"] * penalty
            if next_cost >= best.get(neighbor, math.inf):
                continue
            best[neighbor] = next_cost
            parent[neighbor] = (node, edge_index)
            heuristic = max(0, haversine(nodes[neighbor], goal_coord) - max_goal_offset)
            estimate = next_cost + heuristic
            heapq.heappush(queue, (estimate, next_cost, neighbor))
    if selected_goal is None:
        return None
    path = []
    node = selected_goal
    while node not in start_set:
        previous, edge_index = parent[node]
        path.append((previous, node, edge_index))
        node = previous
    path.reverse()
    return path


def oriented_edge_coords(edge, start_node):
    return edge["coords"] if edge["u"] == start_node else list(reversed(edge["coords"]))


def json_value(value):
    if hasattr(value, "isoformat"):
        return value.isoformat()
    return value


def match_trips(network, trips, stations):
    grid_size = 0.1
    node_grid = defaultdict(list)
    for index, (lon, lat) in enumerate(network["nodes"]):
        node_grid[(math.floor(lon / grid_size), math.floor(lat / grid_size))].append(index)

    def candidate_nodes(coord):
        cell_x, cell_y = math.floor(coord[0] / grid_size), math.floor(coord[1] / grid_size)
        candidates = []
        for radius in range(8):
            for dx in range(-radius, radius + 1):
                for dy in range(-radius, radius + 1):
                    if radius and max(abs(dx), abs(dy)) != radius:
                        continue
                    candidates.extend(node_grid.get((cell_x + dx, cell_y + dy), ()))
            if candidates and radius >= 2:
                break
        if not candidates:
            raise ValueError(f"No railway graph node near {coord}")
        ranked = sorted(candidates, key=lambda index: haversine(coord, network["nodes"][index]))
        nearest_distance = haversine(coord, network["nodes"][ranked[0]])
        connected = [
            index for index in ranked
            if network["adjacency"][index]
            and haversine(coord, network["nodes"][index]) <= nearest_distance + 1.5
        ]
        return (connected or ranked)[:12]

    station_nodes = {}
    station_offsets = {}
    required_stations = {
        name
        for trip in trips
        for name in (*trip["stations"], trip["from"], trip["to"])
        if name in stations
    }
    for name in required_stations:
        coord = stations[name]["coord"]
        candidates = candidate_nodes(coord)
        station_nodes[name] = candidates
        station_offsets[name] = min(haversine(coord, network["nodes"][index]) for index in candidates)

    path_cache = {}
    features = []
    line_edges = defaultdict(set)
    line_counts = Counter()
    line_labels = defaultdict(Counter)
    line_sections = defaultdict(Counter)
    failed_segments = []
    straight_segments = []
    quality_warnings = []

    for position, trip in enumerate(trips, start=1):
        names = [name for name in trip["stations"] if name in station_nodes]
        if len(names) < 2:
            names = [name for name in (trip["from"], trip["to"]) if name in station_nodes]
        route_segments = []
        trip_edge_ids = []
        matched_segments = 0

        for origin, destination in zip(names, names[1:]):
            if origin == destination:
                continue
            key = (origin, destination)
            reverse_key = (destination, origin)
            if key in path_cache:
                path = path_cache[key]
            elif reverse_key in path_cache and path_cache[reverse_key] is not None:
                path = [(v, u, edge_index) for u, v, edge_index in reversed(path_cache[reverse_key])]
            else:
                path = astar(
                    network,
                    station_nodes[origin],
                    station_nodes[destination],
                    stations[origin]["coord"],
                    stations[destination]["coord"],
                )
                path_cache[key] = path

            if path is None:
                failed_segments.append({"trip_id": trip["trip_id"], "from": origin, "to": destination})
                route_segments.append([stations[origin]["coord"], stations[destination]["coord"]])
                straight_segments.append({"trip_id": trip["trip_id"], "from": origin, "to": destination, "reason": "disconnected"})
                continue

            segment_coords = []
            segment_line_keys = set()
            segment_edge_ids = []
            for start_node, _, edge_index in path:
                coords = oriented_edge_coords(network["edges"][edge_index], start_node)
                segment_coords.extend(coords if not segment_coords else coords[1:])
                segment_edge_ids.append(edge_index)
                if network["edges"][edge_index]["line_key"]:
                    segment_line_keys.add(network["edges"][edge_index]["line_key"])
            if len(segment_coords) >= 2:
                direct_distance = haversine(stations[origin]["coord"], stations[destination]["coord"])
                routed_distance = geometry_length(segment_coords)
                if direct_distance and routed_distance / direct_distance > 2.2:
                    route_segments.append([stations[origin]["coord"], stations[destination]["coord"]])
                    straight_segments.append({
                        "trip_id": trip["trip_id"],
                        "from": origin,
                        "to": destination,
                        "reason": "implausible_detour",
                        "detour_ratio": round(routed_distance / direct_distance, 2),
                    })
                    continue
                matched_segments += 1
                route_segments.append(segment_coords)
                trip_edge_ids.extend(segment_edge_ids)
                section = " ↔ ".join(sorted((origin, destination)))
                for line_key in segment_line_keys:
                    line_sections[line_key][section] += 1

        stitched_route = stitch_route_segments(route_segments)
        stated_distance = float(trip["distance"] or 0)
        matched_distance = geometry_length(stitched_route)
        distance_ratio = matched_distance / stated_distance if stated_distance else None
        if distance_ratio is not None and not 0.6 <= distance_ratio <= 1.5:
            quality_warnings.append({
                "trip_id": trip["trip_id"],
                "train": trip["train"],
                "distance_ratio": round(distance_ratio, 3) if distance_ratio is not None else None,
            })
        if len(stitched_route) < 2:
            continue
        trip_lines = defaultdict(set)
        for edge_index in set(trip_edge_ids):
            edge_name = network["edges"][edge_index]["name"].strip()
            line_key = normalized_line_name(edge_name)
            if edge_name and line_key:
                trip_lines[line_key].add(edge_index)
                line_labels[line_key][edge_name] += network["edges"][edge_index]["length_km"]
        line_source = "osm"
        if not trip_lines:
            trip_lines["unclassified"].update(trip_edge_ids)
            line_labels["unclassified"]["未分类线路"] += 1
            line_source = "unclassified"
        for line_key, edge_ids in trip_lines.items():
            line_counts[line_key] += 1
            line_edges[line_key].update(edge_ids)
        total_segments = max(1, len(names) - 1)
        properties = {key: json_value(value) for key, value in trip.items() if key != "stations"}
        properties["fromprov"] = stations.get(trip["from"], {}).get("province", "")
        properties["toprov"] = stations.get(trip["to"], {}).get("province", "")
        properties.update({
            "id": trip["train"],
            "lines": sorted(line_labels[key].most_common(1)[0][0] for key in trip_lines),
            "line_source": line_source,
            "geometry_source": "osm" if matched_segments == total_segments else "mixed",
            "distance_ratio": round(distance_ratio, 3) if distance_ratio is not None else None,
            "match_confidence": round(matched_segments / total_segments, 3),
            "edge_count": len(set(trip_edge_ids)),
        })
        geometry = {"type": "LineString", "coordinates": stitched_route}
        features.append({
            "type": "Feature",
            "properties": properties,
            "geometry": geometry,
        })
        if position % 25 == 0:
            print(f"Matched {position}/{len(trips)} trips")

    network_line_km = defaultdict(float)
    for edge in network["edges"]:
        if edge["line_key"]:
            network_line_km[edge["line_key"]] += edge["length_km"]

    line_stats = []
    for line_key, count in line_counts.most_common():
        edge_ids = line_edges[line_key]
        unique_km = sum(network["edges"][index]["length_km"] for index in edge_ids)
        label = line_labels[line_key].most_common(1)[0][0]
        total_named_km = network_line_km.get(line_key, 0)
        top_section, top_section_trips = line_sections[line_key].most_common(1)[0] if line_sections[line_key] else ("-", 0)
        line_stats.append({
            "line": label,
            "line_key": line_key,
            "trips": count,
            "trip_share": round(count / len(trips), 4) if trips else 0,
            "visited_unique_km": round(unique_km, 1),
            "total_named_track_km": round(total_named_km, 1),
            "coverage": round(min(1, unique_km / total_named_km), 4) if total_named_km else None,
            "visited_edge_count": len(edge_ids),
            "top_section": top_section,
            "top_section_trips": top_section_trips,
        })

    major_lines = [item for item in line_stats if is_major_line(item["line"], item["visited_unique_km"])]

    return features, {
        "summary": {
            "trip_count": len(trips),
            "mapped_trip_count": len(features),
            "classified_trip_count": len(features) - line_counts.get("unclassified", 0),
            "line_count": len([line for line in line_counts if line != "unclassified"]),
            "major_line_count": len(major_lines),
            "failed_segment_count": len(failed_segments),
            "quality_warning_count": len(quality_warnings),
            "straight_segment_count": len(straight_segments),
        },
        "major_lines": major_lines,
        "lines": line_stats,
        "failed_segments": failed_segments,
        "straight_segments": straight_segments,
        "quality_warnings": quality_warnings,
        "station_snap_warnings": [
            {"station": name, "offset_km": round(offset, 2)}
            for name, offset in station_offsets.items() if offset > 3
        ],
    }


def main():
    args = parse_args()
    railway_dir = args.database / "railway_route"
    source = railway_dir / "railways.geojson"
    cache = railway_dir / "rail_network_v3.pkl"
    network = load_network(source, cache, args.rebuild_network)
    stations = load_stations(REFERENCE_DATA / "stations_data.csv")
    trips = load_trips(args.database / "火车乘坐记录.xlsx")
    features, stats = match_trips(network, trips, stations)

    args.output.mkdir(parents=True, exist_ok=True)
    route_path = args.output / "rail_routes_real.geojson"
    stats_path = args.output / "rail_line_stats.json"
    route_path.write_text(json.dumps({"type": "FeatureCollection", "features": features}, ensure_ascii=False, separators=(",", ":")), encoding="utf-8")
    stats_path.write_text(json.dumps(stats, ensure_ascii=False, indent=2), encoding="utf-8")
    print(f"Wrote {len(features)} matched routes to {route_path}")
    print(f"Wrote line statistics to {stats_path}")


if __name__ == "__main__":
    main()
